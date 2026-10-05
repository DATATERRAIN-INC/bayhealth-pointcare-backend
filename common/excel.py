import csv
import io
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from apps.ai_caller.constants import PATIENT_EXCEL_COLUMNS

HEADER_ALIASES = {
    "first_name": "first_name",
    "firstname": "first_name",
    "memberfirstname": "first_name",
    "member_first_name": "first_name",
    "last_name": "last_name",
    "lastname": "last_name",
    "memberlastname": "last_name",
    "member_last_name": "last_name",
    "name": "name",
    "patient_name": "name",
    "address": "address",
    "memberaddress1": "address_line1",
    "member_address1": "address_line1",
    "memberaddress2": "address_line2",
    "member_address2": "address_line2",
    "membercity": "address_city",
    "member_city": "address_city",
    "memberstate": "address_state",
    "member_state": "address_state",
    "memberzipcode": "address_zip",
    "member_zipcode": "address_zip",
    "member_zip": "address_zip",
    "dob": "dob",
    "date_of_birth": "dob",
    "memberdob": "dob",
    "member_dob": "dob",
    "doctor": "doctor",
    "pcpname": "doctor",
    "pcp_name": "doctor",
    "service_name": "service_name",
    "service": "service_name",
    "reason": "service_name",
    "country_code": "country_code",
    "phone": "phone_number",
    "phone_number": "phone_number",
    "mobilephone": "phone_number",
    "mobile_phone": "phone_number",
}


def _normalize_header(value):
    return str(value or "").strip().lower().replace(" ", "_")


def build_patient_template_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Patients"
    sheet.append(list(PATIENT_EXCEL_COLUMNS))
    sheet.append(
        [
            "Maria",
            "Santos",
            "14 Oak Street, Dover, DE 19901",
            "1984-03-12",
            "Dr. Alan Brooks",
            "annual wellness visit",
            "",
            "3025550101",
        ]
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _map_headers(raw_headers):
    mapping = {}
    for index, header in enumerate(raw_headers):
        key = HEADER_ALIASES.get(_normalize_header(header))
        if key and key not in mapping:
            mapping[key] = index
    return mapping


def parse_patient_dob(value):
    """
    Accept common DOB inputs and return a date, or None if invalid.
    Handles datetime/date, Excel serial numbers, and many string formats.
    """
    from datetime import timedelta

    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    # Excel serial date (days since 1899-12-30).
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        serial = float(value)
        # Large numbers are phones, not Excel date serials.
        if serial > 100000:
            text = str(int(serial)) if float(serial).is_integer() else str(value)
            return parse_patient_dob(text)
        try:
            return (datetime(1899, 12, 30) + timedelta(days=int(serial))).date()
        except Exception:
            return None

    original = str(value).strip()
    if not original:
        return None

    # Strip time portion if present.
    stripped = original
    if "T" in stripped:
        stripped = stripped.split("T", 1)[0].strip()
    elif " " in stripped:
        # "1980-11-09 00:00:00" / "Nov 9 1980"
        maybe_date, maybe_time = stripped.split(" ", 1)
        if maybe_time[:1].isdigit() or maybe_time.upper().startswith("AM") or maybe_time.upper().startswith("PM"):
            stripped = maybe_date.strip()

    formats = (
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%m/%d/%Y",
        "%m/%d/%y",
        "%d/%m/%Y",
        "%d/%m/%y",
        "%m-%d-%Y",
        "%m-%d-%y",
        "%d-%m-%Y",
        "%d-%m-%y",
        "%Y%m%d",
        "%m%d%Y",
        "%d%m%Y",
        "%b %d %Y",
        "%B %d %Y",
        "%d %b %Y",
        "%d %B %Y",
        "%b %d, %Y",
        "%B %d, %Y",
        "%d-%b-%Y",
        "%d-%B-%Y",
        "%b-%d-%Y",
        "%B-%d-%Y",
    )

    candidates = []
    for item in (stripped, original.split(" ")[0], original.replace(",", "")):
        item = (item or "").strip()
        if item and item not in candidates:
            candidates.append(item)
        slashy = item.replace("-", "/")
        if slashy and slashy not in candidates:
            candidates.append(slashy)
        dashy = item.replace("/", "-")
        if dashy and dashy not in candidates:
            candidates.append(dashy)

    parsed = None
    for candidate in candidates:
        for fmt in formats:
            try:
                parsed = datetime.strptime(candidate, fmt).date()
                break
            except ValueError:
                continue
        if parsed is not None:
            break

    if parsed is None:
        digits = "".join(ch for ch in stripped if ch.isdigit())
        if len(digits) == 8:
            for fmt in ("%Y%m%d", "%m%d%Y", "%d%m%Y"):
                try:
                    parsed = datetime.strptime(digits, fmt).date()
                    break
                except ValueError:
                    continue

    if parsed is None:
        return None

    today = date.today()
    if parsed > today or parsed.year < 1900:
        return None
    return parsed


def _parse_dob_text(text):
    """Return YYYY-MM-DD string for storage/serializers, or '' if invalid/blank."""
    parsed = parse_patient_dob(text)
    return parsed.isoformat() if parsed else ""


def _cell_to_str(value, *, field=""):
    if value is None:
        return ""
    if field == "dob":
        parsed = parse_patient_dob(value)
        return parsed.isoformat() if parsed else str(value).strip()
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    # Phones may arrive as floats/ints from Excel (e.g. 6688445566.0).
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int) and field == "phone_number":
        return str(value)
    text = str(value).strip()
    if field == "phone_number" and text.endswith(".0"):
        text = text[:-2]
    return text


def _compose_address(row):
    address = (row.get("address") or "").strip()
    if address:
        return address
    line1 = (row.get("address_line1") or "").strip()
    line2 = (row.get("address_line2") or "").strip()
    city = (row.get("address_city") or "").strip()
    state = (row.get("address_state") or "").strip()
    zip_code = (row.get("address_zip") or "").strip()
    city_state = ", ".join(part for part in (city, state) if part)
    if zip_code:
        city_state = f"{city_state} {zip_code}".strip()
    parts = [part for part in (line1, line2, city_state) if part]
    return ", ".join(parts)


def _row_to_dict(values, mapping):
    # Include legacy "name" so services can split it when first/last are blank.
    fields = list(PATIENT_EXCEL_COLUMNS) + [
        "name",
        "address_line1",
        "address_line2",
        "address_city",
        "address_state",
        "address_zip",
    ]
    row = {}
    for field in fields:
        index = mapping.get(field)
        if index is None or index >= len(values):
            row[field] = ""
            continue
        row[field] = _cell_to_str(values[index], field=field)
    row["address"] = _compose_address(row)
    return row

def _iter_xlsx_rows(file_bytes):
    workbook = load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    mapping = _map_headers(rows[0])
    parsed = []
    for row_number, values in enumerate(rows[1:], start=2):
        if values is None or all(cell is None or str(cell).strip() == "" for cell in values):
            continue
        parsed.append((row_number, _row_to_dict(list(values), mapping)))
    return parsed


def _iter_csv_rows(file_bytes):
    text = file_bytes.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return []

    mapping = _map_headers(rows[0])
    parsed = []
    for row_number, values in enumerate(rows[1:], start=2):
        if not values or all(str(cell).strip() == "" for cell in values):
            continue
        parsed.append((row_number, _row_to_dict(values, mapping)))
    return parsed


def parse_patient_upload(file_bytes, filename):
    extension = Path(filename).suffix.lower()
    if extension == ".xlsx":
        return _iter_xlsx_rows(file_bytes)
    if extension == ".csv":
        return _iter_csv_rows(file_bytes)
    raise ValueError("Unsupported file type. Use .xlsx or .csv.")
